"""DOCX для практической работы №7 (вариант 5)."""

from __future__ import annotations

from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _add_figure(doc: Document, image_path: Path | None, caption: str, width_inches: float = 4.2) -> None:
    if image_path is not None and image_path.is_file():
        doc.add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True


def _add_console_placeholder(doc: Document, title: str, instruction: str) -> None:
    doc.add_heading(title, level=3)
    _p(doc, instruction)
    doc.add_paragraph()
    doc.add_paragraph()


def _matrix_table(doc: Document, title: str, mat: np.ndarray, fmt: str = "{:.4g}") -> None:
    doc.add_heading(title, level=3)
    table = doc.add_table(rows=int(mat.shape[0]), cols=int(mat.shape[1]))
    for i in range(mat.shape[0]):
        for j in range(mat.shape[1]):
            table.rows[i].cells[j].text = fmt.format(mat[i, j])


def build_practice07_docx(
    report_path: Path,
    data: dict,
    plot_task2_image: Path | None = None,
    plot_task2_kernel: Path | None = None,
    plot_task2_output: Path | None = None,
    plot_task3_E: Path | None = None,
    source_file_hint: str = "neuro_math_Yakimov_var5/task_7/practice07.py",
) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    t1 = data["task1"]
    t2 = data["task2"]
    t3 = data["task3"]

    doc.add_heading("Практическая работа №7. Вариант 5", level=1)
    _p(
        doc,
        "Работа включает три блока: регрессию целевого признака по бинарным изображениям "
        "(таблица 7.7 методички), аналитическую деконволюцию по данным таблицы 7.8 "
        "и разбор embedding-слоя по рисунку 7.42 с опорой на таблицу 7.9.",
    )

    doc.add_heading("Задание 1. Таблица 7.7 — регрессия Y по изображению", level=2)
    _p(doc, "1) Постановка.")
    _p(
        doc,
        "Для каждого объекта тренировочной выборки задано бинарное изображение и количественное значение Y. "
        "Требуется построить модель, позволяющую оценивать Y по новому изображению того же формата.",
    )
    _p(doc, "2) Реализация в программе.")
    _p(
        doc,
        t1["description"]
        + " В отчётном скрипте использован учебный набор изображений 3×3 и MLPRegressor; "
        "при наличии полного текста таблицы 7.7 в основном пособии значения следует перенести в код.",
    )
    _p(
        doc,
        f"Для демонстрации удержан объект с номером {t1['hold_index_1based']} (прогноз по модели, "
        f"обученной на остальных). Эталонное Y: {t1['y_hold_true']:.6g}, прогноз: {t1['y_hold_pred']:.6g}, "
        f"абсолютная ошибка: {t1['abs_err_hold']:.6g}.",
    )
    _p(
        doc,
        f"На обучающей подвыборке без удержанного объекта: MSE={t1['metrics_train']['mse']:.6g}, "
        f"MAE={t1['metrics_train']['mae']:.6g}, R²={t1['metrics_train']['r2']:.6g}.",
    )

    doc.add_heading("Задание 2. Таблица 7.8 — деконволюция 6×6", level=2)
    _p(doc, "1) Условия.")
    _p(
        doc,
        "Stride по осям x и y равен 1, padding равен 0. Выполняется операция транспонированной свёртки "
        "(деконволюция в терминологии курса), размер выходной карты 6×6.",
    )
    _p(doc, "2) Исходные матрицы (вариант 5).")
    _matrix_table(doc, "Изображение 4×4", t2["image"])
    _matrix_table(doc, "Фильтр 3×3", t2["kernel"])
    _p(doc, "3) Результат.")
    _matrix_table(doc, "Выход 6×6", t2["output_6x6"], fmt="{:.6g}")

    _add_figure(doc, plot_task2_image, "Рисунок 1 — изображение 4×4.")
    _add_figure(doc, plot_task2_kernel, "Рисунок 2 — фильтр 3×3.")
    _add_figure(doc, plot_task2_output, "Рисунок 3 — результат деконволюции 6×6.")

    doc.add_heading("Задание 3. Embedding (рис. 7.42) и таблица 7.9", level=2)
    _p(doc, "1) Модель.")
    _p(
        doc,
        "Задана матрица вложений E размером V×d (число токенов × размерность эмбеддинга). "
        "Для объекта по последовательности индексов токенов вычисляется сумма соответствующих строк E; "
        "скалярный отклик представлен как линейная форма от этой суммы с вектором весов w и смещением b. "
        "В примере параметры w и b подобраны методом наименьших квадратов по учебной таблице.",
    )
    _p(doc, t3["description"])
    _p(
        doc,
        f"Размер словаря V={t3['vocab_size']}, размерность эмбеддинга d={t3['embed_dim']}. "
        f"MSE восстановления Y по таблице: {t3['mse']:.6g}.",
    )
    _add_figure(doc, plot_task3_E, "Рисунок 4 — тепловая карта матрицы E.")

    _add_console_placeholder(
        doc,
        "Рисунок 5 — Скрин консоли (вставить вручную)",
        "Вставьте скрин терминала после `python -m neuro_math_Yakimov_var5.task_7.practice07` "
        "с блоками [1/3]–[3/3] и численными матрицами задания 2.",
    )

    doc.add_heading("Вывод", level=2)
    _p(
        doc,
        "Выполнены расчёты по трём направлениям: регрессия по изображению, деконволюция с заданными "
        "stride и padding, анализ embedding-слоя с линейным выходом по сумме эмбеддингов токенов.",
    )

    doc.add_heading("Ссылка на исходный код", level=2)
    _p(doc, f"Рабочий скрипт: `{source_file_hint}`.")

    doc.save(report_path)
