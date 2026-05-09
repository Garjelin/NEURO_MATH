"""DOCX для практической работы №4 (вариант 5)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _code_block(doc: Document, lines: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(lines.strip("\n"))
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    p.paragraph_format.left_indent = Inches(0.25)


def _add_figure(doc: Document, image_path: Path | None, caption: str, width_inches: float = 5.8) -> None:
    if image_path is not None and image_path.is_file():
        doc.add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True


def build_practice04_docx(
    report_path: Path,
    data: dict,
    plot_task1_tree: Path | None = None,
    plot_task1_pred: Path | None = None,
    plot_task2_tree: Path | None = None,
    plot_task2_bar: Path | None = None,
    source_file_hint: str = "neuro_math_Yakimov_var5/task_4/practice04.py",
) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    t1 = data["task1"]
    t2 = data["task2"]
    t3 = data["task3"]

    doc.add_heading("Практическая работа №4. Вариант 5", level=1)
    _p(
        doc,
        "Работа выполняется по методическому указанию к практическим занятиям №4. "
        "В первых двух заданиях строится дерево решений в постановке регрессии; "
        "в третьем задании задаётся вычислительный граф для функции варианта 5 "
        "и выполняются прямой проход (значение функции и промежуточные узлы) "
        "и обратный проход (градиенты по входам x и y).",
    )

    doc.add_heading("Задание 1. Дерево решений по таблице 4.3", level=2)
    _p(
        doc,
        "Использована полная таблица зависимости (аналог таблицы из методички): десять строк "
        "(x, Y). Для варианта 5 в обучающую выборку включены строки с номерами 7, 8, 9 и 10; "
        "остальные строки образуют логический тестовый набор. В качестве модели применён "
        "DecisionTreeRegressor из scikit-learn с критерием squared_error (стандартная "
        "минимизация MSE при построении дерева). На тесте качество оценивается по среднеквадратичной ошибке MSE.",
    )
    _p(
        doc,
        f"Строки обучения (нумерация как в задании): {t1['train_rows']}. "
        f"Строки теста: {t1['test_rows']}.",
    )
    _p(
        doc,
        f"MSE на обучении: {t1['mse_train']:.6g}. MSE на тесте: {t1['mse_test']:.6g}. "
        "На обучающей части ошибка может быть ненулевой, если одному значению признака соответствуют "
        "разные целевые значения — дерево с одним входом не может одновременно точно аппроксимировать оба.",
    )

    doc.add_heading("Задание 2. Дерево по двум бинарным признакам (таблица 4.5)", level=2)
    _p(
        doc,
        "Признаки x и z принимают значения 0 или 1; целевая переменная Y задана в таблице из методички. "
        "Для варианта 5 обучающая выборка составлена из строк с номерами 2, 3, 5, 8, 9 и 10; "
        "на оставшихся строках оценивается средняя абсолютная ошибка MAE.",
    )
    _p(
        doc,
        f"Критерий разбиения в sklearn: {t2['tree_criterion']}. "
        "При наличии поддержки используется absolute_error, чтобы соответствовать постановке "
        "минимизации MAE при обучении дерева.",
    )
    _p(
        doc,
        f"MAE на обучении: {t2['mae_train']:.6g}. MAE на тесте: {t2['mae_test']:.6g}. "
        "Несколько обучающих строк с одинаковой парой (x, z), но разными Y создают неоднозначность: "
        "дерево выдаёт усреднённое в листе предсказание, что отражается на метриках.",
    )

    doc.add_heading("Задание 3. Граф вычислений и обратное распространение", level=2)
    _p(
        doc,
        "Для варианта 5 задана функция "
        "f(x, y) = exp(x·y) − cos(x/y) + 1/sin(x+y). "
        "Область определения: y ≠ 0 (узел деления x/y) и sin(x+y) ≠ 0 (узел 1/sin(x+y)).",
    )
    _p(doc, "Граф в явном виде (промежуточные узлы):")
    _code_block(
        doc,
        "n1 = x * y\n"
        "n2 = exp(n1)\n"
        "n3 = x / y\n"
        "n4 = cos(n3)\n"
        "n5 = x + y\n"
        "n6 = sin(n5)\n"
        "n7 = 1 / n6\n"
        "f  = n2 - n4 + n7",
    )
    _p(
        doc,
        "Прямой проход вычисляет значения узлов n1…n7 и итоговое f. "
        "Обратный проход реализован вручную по правилу цепочки: от выхода к входам накапливаются "
        "частные производные по промежуточным узлам и затем по x и y. "
        "Для контроля правильности градиенты сравниваются с симметричной конечно-разностной "
        "аппроксимацией в тех же точках.",
    )

    doc.add_heading("Таблица. Примеры точек и градиенты", level=3)
    table = doc.add_table(rows=1, cols=6)
    hdr = table.rows[0].cells
    hdr[0].text = "x"
    hdr[1].text = "y"
    hdr[2].text = "f"
    hdr[3].text = "∂f/∂x (анал.)"
    hdr[4].text = "∂f/∂y (анал.)"
    hdr[5].text = "Контроль (числ.)"
    for row in t3["points"]:
        x, y = row["point"]
        cells = table.add_row().cells
        cells[0].text = f"{x:g}"
        cells[1].text = f"{y:g}"
        cells[2].text = f"{row['f']:.6f}"
        cells[3].text = f"{row['df_dx_analytic']:.6f}"
        cells[4].text = f"{row['df_dy_analytic']:.6f}"
        cells[5].text = f"Δx={row['df_dx_numeric']:.4f}, Δy={row['df_dy_numeric']:.4f}"

    doc.add_heading("Рисунки", level=2)
    _add_figure(
        doc,
        plot_task1_tree,
        "Рисунок 1 — Структура дерева решений (задание 1).",
    )
    _add_figure(
        doc,
        plot_task1_pred,
        "Рисунок 2 — Обучающие и тестовые точки и ступенчатое предсказание дерева по оси x.",
    )
    _add_figure(
        doc,
        plot_task2_tree,
        "Рисунок 3 — Структура дерева решений (задание 2, два признака).",
    )
    _add_figure(
        doc,
        plot_task2_bar,
        "Рисунок 4 — Фактические и предсказанные Y на тестовых строках задания 2.",
    )

    doc.add_heading("Консольный вывод", level=3)
    _p(
        doc,
        "После запуска рабочего скрипта сохраните скрин терминала с блоками [1/3]–[3/3], "
        "метриками MSE/MAE и трассировкой узлов задания 3 — при необходимости вставьте его в отчёт вручную.",
    )
    doc.add_paragraph()
    doc.add_paragraph()

    doc.add_heading("Выводы", level=2)
    _p(
        doc,
        "Построены регрессионные деревья решений для табличных данных с метриками MSE и MAE "
        "согласно варианту 5. Для функции из таблицы 4.7 записан вычислительный граф, "
        "реализованы прямой и обратный проходы; совпадение аналитического градиента "
        "с численной проверкой подтверждает корректность обратного распространения.",
    )

    doc.add_heading("Ссылка на исходный код", level=2)
    _p(
        doc,
        f"Рабочий код задания: `{source_file_hint}`. "
        "Генератор DOCX находится в `report_generator/` и не является частью основного решения.",
    )

    doc.save(report_path)
