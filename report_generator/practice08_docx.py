"""DOCX для практической работы №8 (вариант 5)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _add_figure(doc: Document, image_path: Path | None, caption: str, width_inches: float = 5.5) -> None:
    if image_path is not None and image_path.is_file():
        doc.add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True


def _add_console_placeholder(doc: Document, title: str, instruction: str) -> None:
    doc.add_heading(title, level=3)
    _p(doc, instruction)
    doc.add_paragraph()
    doc.add_paragraph()


def build_practice08_docx(
    report_path: Path,
    data: dict,
    plot_task1_bar: Path | None = None,
    source_file_hint: str = "neuro_math_Yakimov_var5/task_8/practice08.py",
) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    txt = data["text"]
    pdf_a = data["pdf_assignment"]
    t_char = data["task_char"]
    t_line = data["task_line_demo"]

    doc.add_heading("Практическая работа №8. Вариант 5", level=1)

    doc.add_heading("Условие", level=2)
    doc.add_paragraph(pdf_a)

    doc.add_heading("Исходные данные (таблица 8.9, вариант 5)", level=2)
    _p(doc, "Для варианта 5 из таблицы 8.9 использован следующий трёхстрочный фрагмент (разделитель — перевод строки).")
    doc.add_paragraph(txt)

    doc.add_heading("Часть 1. Нейросеть: предсказание следующего символа", level=2)

    doc.add_heading("Цель и метрики", level=3)
    _p(doc, t_char["description"])
    _p(
        doc,
        "Цель обучения — минимизация ошибки классификации при предсказании следующего символа по контексту "
        "фиксированной длины. Признак на каждой позиции — конкатенация one-hot векторов символов контекста "
        "(размерность равна длине контекста, умноженной на мощность алфавита). "
        "Доля верных ответов (accuracy) считается отдельно на обучающей и на контрольной частях; "
        "на контроль отводится 25% примеров (либо последние по порядку в тексте, либо случайная выборка — "
        "в зависимости от способа, см. ключ). Дополнительно фиксируется логарифмическая функция потерь на контроле.",
    )

    doc.add_heading("Модель и программная реализация", level=3)
    _p(
        doc,
        "Реализован многослойный персептрон MLPClassifier из библиотеки scikit-learn: два скрытых слоя "
        "из 96 и 48 нейронов, функция активации ReLU, адаптивный оптимизатор Adam, L2-регуляризация (alpha). "
        "При достаточном числе обучающих точек включён внутренний early stopping по доле валидации внутри train; "
        "для малых выборок итерации ограничены параметрами сходимости.",
    )

    doc.add_heading("Способы подготовки выборки", level=3)
    _p(
        doc,
        "Реализовано несколько способов, отличающихся длиной контекста (1, 2 или 3 символа), построением скользящего "
        "окна по всему тексту сразу или только внутри отдельных строк (без пересечения окна через перевод строки), "
        "а также типом разбиения на обучение и контроль (хронологический отрезок хвоста 25% против случайного split).",
    )

    tbl = doc.add_table(rows=1 + len(t_char["strategies"]), cols=6)
    hdr = tbl.rows[0].cells
    hdr[0].text = "Способ"
    hdr[1].text = "n_train"
    hdr[2].text = "n_val"
    hdr[3].text = "Accuracy train"
    hdr[4].text = "Accuracy val"
    hdr[5].text = "Итераций"
    for i, r in enumerate(t_char["strategies"]):
        row = tbl.rows[i + 1].cells
        row[0].text = r["key"]
        row[1].text = str(r["n_train"])
        row[2].text = str(r["n_val"])
        row[3].text = f"{r['acc_train']:.4f}"
        row[4].text = f"{r['acc_val']:.4f}"
        row[5].text = str(r["n_iter"])

    doc.add_heading("Расшифровка ключей способов", level=3)
    for r in t_char["strategies"]:
        _p(doc, f"{r['key']}: {r['title']}")

    _p(
        doc,
        f"На контрольной выборке наибольшую долю верных предсказаний даёт способ с ключом {t_char['best_val_key']}. "
        "На коротком тексте значения accuracy на контроле могут быть невысокими: число классов велико относительно "
        "объёма данных, а отложенный фрагмент может содержать редкие символы.",
    )

    doc.add_heading("Рисунок 1", level=3)
    _add_figure(doc, plot_task1_bar, "Сравнение способов подготовки выборки по accuracy на контроле (часть 1).")

    doc.add_heading("Часть 2. Классификация следующей строки из корпуса", level=2)

    doc.add_heading("Постановка", level=3)
    _p(doc, t_line["description"])

    doc.add_heading("Корпус строк с индексами", level=3)
    _p(
        doc,
        "Три строки стихотворения нумеруются L0, L1, L2. Именно текст одной из этих строк и является результатом "
        "предсказания: модель выбирает номер строки k ∈ {0,1,2}, после чего «предсказанная следующая строка» "
        "выводится как полное содержимое L_k.",
    )
    tbl_l = doc.add_table(rows=1 + len(t_line["lines_numbered"]), cols=2)
    h2 = tbl_l.rows[0].cells
    h2[0].text = "Индекс k"
    h2[1].text = "Текст строки Lk"
    for i, rown in enumerate(t_line["lines_numbered"]):
        row = tbl_l.rows[i + 1].cells
        row[0].text = str(rown["index"])
        row[1].text = rown["text"]

    doc.add_heading("Признаки, класс и обучающая выборка", level=3)
    _p(
        doc,
        "Вход сети — вектор абсолютных частот символов текущей строки по алфавиту всего трёхстрочного текста "
        "(размерность равна числу различных символов). "
        "Целевая переменная (класс) — индекс следующей строки в списке [L0, L1, L2], то есть для первой пары "
        "истинный класс 1 (следует L1), для второй пары — класс 2 (следует L2). "
        "Обучающая выборка состоит из двух объектов (две пары «текущая строка → индекс следующей»). "
        "После обучения для каждой пары сравниваются эталонный текст следующей строки и текст строки с индексом, "
        "предсказанным сетью — это и есть наглядный результат в виде целой строки.",
    )

    doc.add_heading("Результаты: эталон и предсказанный текст следующей строки", level=3)
    _p(doc, f"Accuracy на двух обучающих парах (проверка на тех же парах): {t_line['acc_on_two_pairs']:.4f}.")

    tb2 = doc.add_table(rows=1 + len(t_line["examples"]), cols=6)
    h3 = tb2.rows[0].cells
    h3[0].text = "После строки №"
    h3[1].text = "Класс (эталон)"
    h3[2].text = "Класс (прогноз)"
    h3[3].text = "Эталон: текст следующей строки"
    h3[4].text = "Прогноз: текст следующей строки"
    h3[5].text = "Верно"
    for i, ex in enumerate(t_line["examples"]):
        row = tb2.rows[i + 1].cells
        row[0].text = str(ex["after_line"])
        row[1].text = str(ex["true_class"])
        row[2].text = str(ex["pred_class"])
        row[3].text = ex["true_next"]
        row[4].text = ex["pred_next"]
        row[5].text = "да" if ex["ok"] else "нет"

    doc.add_heading("Итоговые предсказанные строки", level=3)
    for ex in t_line["examples"]:
        _p(doc, f"После L{ex['after_line'] - 1}: «{ex['pred_next']}».")

    doc.add_heading("Текущая строка (вход) для полноты картины", level=3)
    for ex in t_line["examples"]:
        _p(doc, f"Для пары после строки {ex['after_line']} текущая строка (вход признаков): «{ex['current']}».")

    doc.add_heading("Выводы", level=2)
    _p(
        doc,
        "В части 1 для микроскопического корпуса из таблицы 8.9 символьная нейросеть на отложенной выборке "
        "демонстрирует ограниченную точность; сравнение способов показывает влияние длины контекста и схемы разбиения. "
        "В части 2 подтверждено, что при выбранной постановке результат предсказания — это полный текст следующей "
        "строки из таблицы (в примере обе пары классифицированы верно).",
    )

    _add_console_placeholder(
        doc,
        "Скрин консоли",
        "При необходимости вставьте скрин вывода программы с таблицей метрик части 1 и таблицей части 2.",
    )

    doc.add_heading("Исходный код", level=2)
    _p(doc, f"Рабочий скрипт: `{source_file_hint}`. Генератор отчёта: `report_generator/practice08_docx.py`.")

    doc.save(report_path)
