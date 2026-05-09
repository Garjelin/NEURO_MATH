"""Развёрнутый DOCX для практической работы №3 (вариант 5)."""

from __future__ import annotations

from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _code_block(doc: Document, lines: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(lines.strip("\n"))
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    p.paragraph_format.left_indent = Inches(0.25)


def _add_figure(doc: Document, image_path: Path | None, caption: str, width_inches: float = 5.8) -> None:
    if image_path is not None and image_path.is_file():
        doc.add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True


def _table_metrics(doc: Document, title: str, z_metrics: dict, mm_metrics: dict) -> None:
    doc.add_heading(title, level=3)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Метод нормализации"
    hdr[1].text = "MSE"
    hdr[2].text = "MAE"
    hdr[3].text = "R²"

    row = table.add_row().cells
    row[0].text = "Z-score"
    row[1].text = f"{z_metrics['mse']:.6f}"
    row[2].text = f"{z_metrics['mae']:.6f}"
    row[3].text = f"{z_metrics['r2']:.6f}"

    row = table.add_row().cells
    row[0].text = "Min-Max"
    row[1].text = f"{mm_metrics['mse']:.6f}"
    row[2].text = f"{mm_metrics['mae']:.6f}"
    row[3].text = f"{mm_metrics['r2']:.6f}"


def build_practice03_docx(
    report_path: Path,
    data: dict,
    plot_scaling: Path | None = None,
    plot_test: Path | None = None,
    source_file_hint: str = "neuro_math_Yakimov_var5/task_3/practice03.py",
) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    split = data["split"]
    z = data["zscore"]
    mm = data["minmax"]

    doc.add_heading("Практическая работа №3. Вариант 5", level=1)
    _p(
        doc,
        "Тема задания по таблице 3.11: «Прогноз стоимости акций». "
        "По условию необходимо выполнить нормализацию данных двумя способами, "
        "самостоятельно сформировать тренировочную выборку, обучить нейронную сеть "
        "и сравнить результат на тестовой выборке.",
    )
    _p(
        doc,
        "В работе используется единая модель регрессии MLPRegressor и единая выборка признаков; "
        "изменяется только способ нормализации: (1) стандартизация Z-score и (2) масштабирование Min-Max.",
    )

    doc.add_heading("1-й этап: подготовка данных", level=2)
    _p(
        doc,
        "Сформирован воспроизводимый временной ряд (синтетические цены закрытия и объёмы торгов). "
        "Из ряда получены признаки: текущая и лаговые лог-доходности, скользящая волатильность, "
        "объём, текущий уровень цены. Целевая переменная — цена следующего шага.",
    )
    _p(
        doc,
        f"Итоговый объём supervised-выборки: {split['n_total']} наблюдений. "
        f"Разбиение по времени: {split['n_train']} обучающих и {split['n_total'] - split['n_train']} тестовых точек.",
    )
    _p(
        doc,
        "Разбиение хронологическое (без shuffle), что корректно для временных рядов "
        "и исключает утечку будущей информации в обучение.",
    )

    doc.add_heading("2-й этап: два способа нормализации", level=2)
    _p(
        doc,
        "Способ 1 — Z-score (StandardScaler): из каждого признака вычитается среднее и "
        "делится на стандартное отклонение. В результате признаки имеют примерно нулевое "
        "среднее и единичную дисперсию.",
    )
    _p(
        doc,
        "Способ 2 — Min-Max (MinMaxScaler): каждый признак линейно переводится в диапазон [0, 1]. "
        "Метод удобен, когда важно ограничить масштаб входа фиксированным интервалом.",
    )
    _p(
        doc,
        "Для корректного сравнения архитектура сети, параметры обучения и разбиение train/test "
        "для обоих способов идентичны. Отличается только предобработка.",
    )

    doc.add_heading("3-й этап: обучение модели", level=2)
    _p(
        doc,
        "Использована модель MLPRegressor (слои 48 и 24 нейрона, активация tanh, решатель L-BFGS). "
        "Целевая переменная также масштабируется соответствующим скейлером, после прогнозирования "
        "предсказания переводятся обратно в исходный масштаб цены.",
    )
    _code_block(
        doc,
        "Запуск рабочего скрипта (без генерации отчёта):\n"
        "cd neuro_math_Yakimov_var5\n"
        "python -m task_3.practice03",
    )

    doc.add_heading("4-й этап: тестирование и сравнение методов нормализации", level=2)
    _p(
        doc,
        "Качество сравнивается по метрикам MSE, MAE и R² на обучающей и тестовой частях. "
        "Ниже приведены значения, полученные автоматическим запуском при формировании отчёта.",
    )

    _table_metrics(doc, "Таблица 1. Метрики на обучающей выборке", z["metrics_train"], mm["metrics_train"])
    _table_metrics(doc, "Таблица 2. Метрики на тестовой выборке", z["metrics_test"], mm["metrics_test"])

    _p(
        doc,
        f"Для теста: RMSE(Z-score) = {float(np.sqrt(z['metrics_test']['mse'])):.6f}, "
        f"RMSE(Min-Max) = {float(np.sqrt(mm['metrics_test']['mse'])):.6f}. "
        "Меньшее RMSE и MAE, а также большее R² указывают на более предпочтительный "
        "способ нормализации для данной постановки.",
    )

    _add_figure(
        doc,
        plot_scaling,
        "Рисунок 1 — Демонстрация эффекта нормализации (исходный признак, Z-score, Min-Max).",
    )
    _add_figure(
        doc,
        plot_test,
        "Рисунок 2 — Сравнение прогнозов на тесте: факт, Z-score, Min-Max.",
    )

    doc.add_heading("Рисунок 3 — скриншот консоли (вставить вручную)", level=3)
    _p(
        doc,
        "После запуска рабочего скрипта сделайте скрин терминала, чтобы были видны блоки [1/4]–[4/4], "
        "метрики обоих способов нормализации и пути к графикам. "
        "Вставьте скриншот в это место через «Вставка → Рисунки».",
    )
    doc.add_paragraph()
    doc.add_paragraph()

    doc.add_heading("5-й этап: выводы", level=2)
    _p(
        doc,
        "Требование задания выполнено: данные нормализованы двумя способами, на каждой нормализации "
        "обучена нейронная сеть и проведено сравнение на тестовой выборке. "
        "Сделан вывод о предпочтительном методе нормализации по метрикам качества.",
    )
    _p(
        doc,
        "Возможные улучшения: применение реальных котировок, walk-forward валидация, "
        "расширение набора финансовых индикаторов и настройка гиперпараметров сети.",
    )

    doc.add_heading("Ссылка на исходный код", level=2)
    _p(
        doc,
        f"Рабочий код задания: `{source_file_hint}`. "
        "Скрипт генерации отчёта находится в `report_generator/` и не является частью "
        "основного исполняемого решения.",
    )

    doc.save(report_path)
