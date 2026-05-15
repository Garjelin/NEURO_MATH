"""DOCX для практической работы №9 (вариант 5)."""

from __future__ import annotations

from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _add_figure(doc: Document, image_path: Path | None, caption: str, width_inches: float = 5.5) -> None:
    if image_path is not None and image_path.is_file():
        doc.add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True


def _add_matrix_table(doc: Document, label: str, mat: np.ndarray, row_labels: list[str], col_labels: list[str]) -> None:
    _p(doc, label)
    table = doc.add_table(rows=1 + mat.shape[0], cols=1 + mat.shape[1])
    table.rows[0].cells[0].text = "Эталон \\ прогноз"
    for j, name in enumerate(col_labels):
        table.rows[0].cells[j + 1].text = name
    for i, name in enumerate(row_labels):
        table.rows[i + 1].cells[0].text = name
        for j in range(mat.shape[1]):
            table.rows[i + 1].cells[j + 1].text = str(int(mat[i, j]))


def _add_console_placeholder(doc: Document) -> None:
    doc.add_heading("Скрин консоли", level=2)
    _p(doc, "При необходимости вставьте скрин вывода `python -m neuro_math_Yakimov_var5.task_9.practice09`.")
    doc.add_paragraph()
    doc.add_paragraph()


def _add_classification_report_table(doc: Document, report: dict, styles: list[str]) -> None:
    table = doc.add_table(rows=1 + len(styles), cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Стиль"
    hdr[1].text = "Precision"
    hdr[2].text = "Recall"
    hdr[3].text = "F1-score"
    for i, name in enumerate(styles):
        row = table.rows[i + 1].cells
        row[0].text = name
        row[1].text = f"{report[name]['precision']:.4f}"
        row[2].text = f"{report[name]['recall']:.4f}"
        row[3].text = f"{report[name]['f1-score']:.4f}"


def build_practice09_docx(
    report_path: Path,
    data: dict,
    plot_samples: Path | None = None,
    plot_confusion: Path | None = None,
    plot_queries: Path | None = None,
    source_file_hint: str = "neuro_math_Yakimov_var5/task_9/practice09.py",
) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    train = data["train"]
    styles = data["styles"]

    doc.add_heading("Практическая работа №9. Вариант 5", level=1)

    doc.add_heading("Условие", level=2)
    doc.add_paragraph(data["pdf_assignment"])

    doc.add_heading("Постановка задачи", level=2)
    _p(
        doc,
        "Для варианта 5 требуется модель подбора художественного стиля изображения. "
        "Формируется воспроизводимый набор изображений с признаками нескольких "
        "стилей WikiArt-подобного корпуса: цветовая палитра, насыщенность, контраст, текстура и геометричность. "
        "Задача решается как многоклассовая классификация: по численным признакам изображения сеть выбирает "
        "один из художественных стилей.",
    )
    _p(
        doc,
        "Входом модели является вектор признаков, вычисленный по изображению. Признаковое описание снижает "
        "размерность входных данных и концентрирует информацию о цвете, контрасте, яркости и структуре изображения. "
        "Далее задача сводится к обучению классификатора, сопоставляющего вектор признаков одному из заданных "
        "художественных стилей.",
    )

    doc.add_heading("Стили и признаки", level=2)
    _p(doc, "Используемые стили: " + ", ".join(styles) + ".")
    _p(
        doc,
        f"Размер набора: {data['n_samples']} изображений; число признаков на одно изображение: {data['n_features']}. "
        "Признаки включают средние значения и стандартные отклонения RGB-каналов, насыщенность, яркость, "
        "энергию границ, блочность и диапазон яркости.",
    )
    _p(
        doc,
        "Каждый стиль задаётся набором цветовых и структурных правил. Для импрессионизма используются короткие "
        "цветовые мазки и мягкая палитра; для кубизма — прямоугольные области и контрастные границы; для "
        "экспрессионизма — насыщенные контрастные цвета и длинные мазки; для реализма — более плавные градиенты "
        "и приглушённые тона; для абстракционизма — смешение ярких геометрических форм и цветовых пятен. "
        "За счёт фиксированного генератора случайных чисел набор воспроизводим: при повторном запуске получаются "
        "те же изображения и те же численные результаты.",
    )
    _p(
        doc,
        "Для каждого изображения рассчитывается 13 признаков: средние RGB-компоненты, стандартные отклонения "
        "RGB-компонент, средняя насыщенность, разброс насыщенности, средняя яркость, разброс яркости, энергия "
        "границ, показатель блочности и интервал между 90-м и 10-м процентилями яркости. Эти признаки позволяют "
        "учесть как цветовое решение, так и фактуру изображения.",
    )
    _add_figure(doc, plot_samples, "Рисунок 1 — примеры изображений набора по стилям.")

    doc.add_heading("Модель", level=2)
    _p(
        doc,
        "Перед обучением признаки стандартизируются с помощью StandardScaler. "
        "В качестве модели используется MLPClassifier: два скрытых слоя (64 и 32 нейрона), "
        "активация ReLU, оптимизатор Adam, L2-регуляризация. Разбиение: 75% обучающая часть и 25% тестовая часть, "
        "с сохранением долей классов.",
    )
    _p(
        doc,
        "Стандартизация признаков необходима, потому что разные характеристики имеют разные диапазоны: средние "
        "значения RGB лежат в интервале 0…1, а показатели контраста и текстуры имеют другую дисперсию. "
        "После стандартизации все признаки оказываются сопоставимыми по масштабу, что упрощает обучение "
        "градиентным методом. Два скрытых слоя позволяют учитывать нелинейные сочетания цвета и структуры.",
    )

    doc.add_heading("Качество модели", level=2)
    _p(
        doc,
        f"n_train = {train['n_train']}, n_test = {train['n_test']}, число итераций обучения = {train['n_iter']}. "
        f"Accuracy на обучении = {train['accuracy_train']:.4f}; accuracy на тесте = {train['accuracy_test']:.4f}.",
    )
    _p(
        doc,
        "Accuracy показывает долю изображений тестовой выборки, для которых выбран правильный стиль. "
        "Дополнительно используется матрица ошибок: строки соответствуют эталонному стилю, столбцы — прогнозу. "
        "Диагональные элементы — верные классификации, внедиагональные элементы — случаи смешения стилей.",
    )
    _add_matrix_table(doc, "Матрица ошибок на тестовой части:", train["confusion"], styles, styles)
    _p(
        doc,
        "По матрице ошибок видно, что импрессионизм, кубизм, экспрессионизм и реализм в данном эксперименте "
        "распознаются без ошибок на тестовой части. Основные смешения возникают для абстракционизма: часть таких "
        "изображений модель относит к импрессионизму и экспрессионизму, так как эти классы имеют пересекающиеся "
        "цветовые и текстурные признаки: яркие цветовые пятна, высокий контраст и выраженную фактуру.",
    )
    _add_figure(doc, plot_confusion, "Рисунок 2 — матрица ошибок по стилям.")

    doc.add_heading("Показатели по каждому стилю", level=3)
    _p(
        doc,
        "Для более детальной оценки приведены precision, recall и F1-score. Precision показывает, насколько "
        "чистыми являются прогнозы данного класса; recall — какую долю объектов данного эталонного класса "
        "модель нашла; F1-score объединяет обе характеристики.",
    )
    _add_classification_report_table(doc, train["report"], styles)

    doc.add_heading("Подбор стиля для контрольных изображений", level=2)
    _p(
        doc,
        "После обучения модель применяется к трём контрольным изображениям, которые не используются при обучении. "
        "Для каждого изображения выводится наиболее вероятный стиль и три стиля с максимальными вероятностями. "
        "Это позволяет оценить не только итоговый выбор, но и степень уверенности модели.",
    )
    table = doc.add_table(rows=1 + len(data["predictions"]), cols=5)
    hdr = table.rows[0].cells
    hdr[0].text = "Изображение"
    hdr[1].text = "Ожидаемый стиль"
    hdr[2].text = "Подобранный стиль"
    hdr[3].text = "Уверенность"
    hdr[4].text = "Top-3 вероятности"
    for i, item in enumerate(data["predictions"]):
        row = table.rows[i + 1].cells
        row[0].text = item["name"]
        row[1].text = item["expected_style"]
        row[2].text = item["predicted_style"]
        row[3].text = f"{item['confidence']:.4f}"
        row[4].text = "; ".join(f"{name}: {prob:.3f}" for name, prob in item["top3"])
    _add_figure(doc, plot_queries, "Рисунок 3 — контрольные изображения и подобранные стили.")
    _p(
        doc,
        "Во всех трёх контрольных примерах выбранный стиль совпал с ожидаемым. Значения вероятностей не равны "
        "единице, так как изображения имеют общие признаки с другими стилями: например, яркие пятна и насыщенные "
        "цвета могут быть характерны как для абстрактной композиции, так и для экспрессионистского изображения. "
        "Поэтому top-3 вероятности полезны для анализа близости стилей.",
    )

    doc.add_heading("Выводы", level=2)
    _p(
        doc,
        "Построена модель, подбирающая художественный стиль изображения по набору численных признаков. "
        "На тестовой части получена высокая accuracy, а матрица ошибок показывает, какие стили различаются "
        "уверенно и где возникают смешения. Для контрольных изображений выведены подобранные стили и вероятности top-3.",
    )
    _p(
        doc,
        "Полученные результаты показывают, что выбранное признаковое пространство позволяет разделять стили "
        "по цветовым и структурным характеристикам. Наиболее устойчиво выделяются классы с явно выраженной "
        "геометрией, контрастом или приглушённой палитрой; смешения возникают между стилями с близкими цветовыми "
        "и фактурными свойствами. Общая схема решения включает подготовку набора данных, выделение признаков, "
        "обучение классификатора, проверку качества и применение модели к новым изображениям.",
    )

    _add_console_placeholder(doc)

    doc.add_heading("Исходный код", level=2)
    _p(doc, f"Рабочий скрипт: `{source_file_hint}`. Генератор отчёта: `report_generator/practice09_docx.py`.")

    doc.save(report_path)
