"""DOCX для практической работы №6 (вариант 5): dropout при прогнозе как в практике 2 зад. 2."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _add_figure(doc: Document, image_path: Path | None, caption: str, width_inches: float = 6.0) -> None:
    if image_path is not None and image_path.is_file():
        doc.add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True


def _add_console_placeholder(doc: Document, title: str, instruction: str) -> None:
    doc.add_heading(title, level=3)
    _p(doc, instruction)
    doc.add_paragraph()
    doc.add_paragraph()


def build_practice06_docx(
    report_path: Path,
    data: dict,
    plot_forecasts: Path | None = None,
    plot_mse_bar: Path | None = None,
    source_file_hint: str = "neuro_math_Yakimov_var5/task_6/practice06.py",
) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    exp = data["experiment"]
    split = exp["split"]
    cfg = exp["training_config"]

    doc.add_heading("Практическая работа №6. Вариант 5", level=1)
    _p(
        doc,
        "Цель работы — исследовать влияние dropout на качество нейросетевого прогноза в постановке, "
        "совпадающей с практической работой №2, задание 2: предсказание цены закрытия на следующий шаг "
        "по набору финансово-технических признаков на синтетическом OHLCV-подобном ряду.",
    )
    _p(
        doc,
        "Важное техническое замечание: модуль `MLPRegressor` из scikit-learn не реализует dropout. "
        "Поэтому для выполнения задания построена эквивалентная полносвязная сеть с той же топологией "
        "скрытых слоёв (32 и 16 нейронов) и активацией tanh, обучаемая на NumPy (мини-батч Adam) "
        "с inverted dropout после каждого скрытого слоя. Так решается задача без установки тяжёлого "
        "пакета PyTorch (`torch`). Вероятность «выключения» нейрона задаётся параметром p.",
    )

    doc.add_heading("1. Постановка и данные", level=2)
    _p(
        doc,
        "Признаки и хронологическое разбиение train/test формируются по тем же правилам, что "
        "в `practice02.run_task2_stock_forecast`: лаги лог-доходности, объём, скользящая волатильность, "
        "текущий уровень цены; цель — значение цены на следующий день.",
    )
    _p(
        doc,
        f"Размер эффективной выборки (после формирования лагов): {split['n_eff']} наблюдений. "
        f"Обучающая часть: {split['n_train']} точек; тестовая — оставшиеся. "
        "Признаки и целевая переменная на обучении стандартизуются через `StandardScaler` так же, как в практике №2.",
    )

    doc.add_heading("2. Эксперимент по dropout", level=2)
    _p(
        doc,
        "Согласно заданию, выполняются три прогона с вероятностью dropout p = 0.1, p = 0.2 и p = 0.5. "
        "При обучении dropout активен (режим train); при получении прогноза на тесте используется режим eval, "
        "при котором dropout выключен.",
    )
    _p(
        doc,
        f"Параметры оптимизации (единые для всех p): число эпох {cfg['epochs']}, размер мини-батча {cfg['batch_size']}, "
        f"скорость обучения Adam {cfg['lr']}, L2-регуляризация (weight_decay) {cfg['weight_decay']}, "
        f"реализация: {cfg['backend']}.",
    )

    doc.add_heading("3. Результаты на тестовой выборке", level=2)
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    hdr[0].text = "Dropout p"
    hdr[1].text = "MSE"
    hdr[2].text = "MAE"
    hdr[3].text = "R²"
    hdr[4].text = "Примечание"
    for run in exp["runs"]:
        mt = run["metrics_test"]
        row = table.add_row().cells
        row[0].text = f"{run['dropout_p']:g}"
        row[1].text = f"{mt['mse']:.6g}"
        row[2].text = f"{mt['mae']:.6g}"
        row[3].text = f"{mt['r2']:.6g}"
        row[4].text = "Тест, хронология сохранена"

    _p(
        doc,
        "Интерпретация: при повышении p регуляризация усиливается, что может уменьшить переобучение, "
        "но при слишком большом p модель теряет ёмкость и качество на тесте может ухудшиться. "
        "Конкретное соотношение метрик зависит от масштаба данных и числа эпох.",
    )

    doc.add_heading("4. Иллюстрации", level=2)
    _add_figure(
        doc,
        plot_forecasts,
        "Рисунок 1 — Тестовый интервал: фактическая цена и прогнозы при разных значениях dropout.",
    )
    _add_figure(
        doc,
        plot_mse_bar,
        "Рисунок 2 — Сравнение среднеквадратичной ошибки на тесте для p ∈ {0.1, 0.2, 0.5}.",
    )

    _add_console_placeholder(
        doc,
        "Рисунок 3 — Скрин консоли (вставить вручную)",
        "Вставьте скрин терминала после запуска `python -m neuro_math_Yakimov_var5.task_6.practice06`: "
        "должны быть видны параметры разбиения выборки, настройки обучения и таблица метрик на тесте для каждого p.",
    )

    doc.add_heading("Вывод", level=2)
    _p(
        doc,
        "Проведено сравнение трёх значений dropout в эквивалентной постановке практики №2 (задание 2). "
        "Зафиксированы метрики регрессии на отложенном тесте и графические зависимости прогноза от выбора p.",
    )

    doc.add_heading("Ссылка на исходный код", level=2)
    _p(doc, f"Рабочий скрипт: `{source_file_hint}`. Генератор отчёта: `report_generator/practice06_docx.py`.")

    doc.save(report_path)
